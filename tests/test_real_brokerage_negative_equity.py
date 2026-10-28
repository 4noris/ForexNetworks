"""Artificial regression fixtures; no broker account records."""
from __future__ import annotations
from datetime import date
import unittest
import pandas as pd
from openpyxl import load_workbook
import api
NEGATIVE_EQUITY_FIXTURE = [(900201, 'Fixture Client 1', -210.0), (900202, 'Fixture Client 2', -125.5), (900203, 'Fixture Client 3', -75.25), (900204, 'Fixture Client 4', -50.0), (900205, 'Fixture Client 5', -12.75), (900206, 'Fixture Client 6', -8.0), (900207, 'Fixture Client 7', -1.25)]
POSITIVE_OPENING_EQ = 1000.0
POSITIVE_CLOSING_EQ = 1100.0
POSITIVE_LOGIN = 200001

def _snapshot_rows(positive_equity: float) -> list[dict[str, object]]:
    rows = [{'Login': login, 'Name': name, 'Equity': equity} for login, name, equity in NEGATIVE_EQUITY_FIXTURE]
    rows.append({'Login': POSITIVE_LOGIN, 'Name': 'Positive Account', 'Equity': positive_equity})
    return rows

def _account_rows(opening_equity: float, closing_equity: float) -> list[dict[str, object]]:
    rows = []
    for login, name, equity in NEGATIVE_EQUITY_FIXTURE:
        rows.append({'Login': login, 'Name': name, 'Group': 'demo', 'Currency': 'USD', 'Closed Lots': 0.0, 'NET DP/WD': 0.0, 'Credit': 0.0, 'Opening Equity': equity, 'Closing Equity': equity, 'NET PNL USD': 0.0, 'NET PNL %': 0.0, 'Deposit': 0.0, 'Withdrawal': 0.0, 'Commission': 0.0, 'Swap': 0.0, 'EOD Closing Equity Date': '2026-05-08'})
    rows.append({'Login': POSITIVE_LOGIN, 'Name': 'Positive Account', 'Group': 'demo', 'Currency': 'USD', 'Closed Lots': 0.0, 'NET DP/WD': 0.0, 'Credit': 0.0, 'Opening Equity': opening_equity, 'Closing Equity': closing_equity, 'NET PNL USD': round(closing_equity - opening_equity, 2), 'NET PNL %': round((closing_equity - opening_equity) / opening_equity * 100.0, 2), 'Deposit': 0.0, 'Withdrawal': 0.0, 'Commission': 0.0, 'Swap': 0.0, 'EOD Closing Equity Date': '2026-05-08'})
    return rows

def _base_payload(clip_negative_equity: bool) -> dict[str, object]:
    return {'clip_negative_equity': clip_negative_equity, 'profile_name': 'Fixture Broker', 'selected_groups': ['demo'], 'summary_range': (date(2026, 5, 2), date(2026, 5, 8)), 'warnings': [], 'opening_report': _snapshot_rows(POSITIVE_OPENING_EQ), 'closing_report': _snapshot_rows(POSITIVE_CLOSING_EQ), 'accounts_display': _account_rows(POSITIVE_OPENING_EQ, POSITIVE_CLOSING_EQ), 'group_summary': [{'Group': 'demo', 'Accounts': len(NEGATIVE_EQUITY_FIXTURE) + 1, 'Closed_Lots': 0.0, 'NET_DP_WD': 0.0, 'Credit': 0.0, 'NET_PNL_USD': 100.0}], 'books_summary': [{'Accounts': len(NEGATIVE_EQUITY_FIXTURE) + 1, 'Closed_Lots': 0.0, 'NET_DP_WD': 0.0, 'Credit': 0.0, 'NET_PNL_USD': 100.0}], 'summary': [], 'summary_with_total': [], 'daily_summary': [], 'lp_summary': [], 'brokerage_totals': {}}

def _load_table_rows(ws, header_label: str='Login') -> list[dict[str, object]]:
    headers: list[object] | None = None
    rows: list[dict[str, object]] = []
    for row in ws.iter_rows(values_only=True):
        if headers is None:
            if row and row[0] == header_label:
                headers = list(row)
            continue
        if row[0] in (None, ''):
            break
        rows.append({str(headers[idx]): row[idx] for idx in range(len(headers)) if headers[idx] is not None})
    return rows

def _load_metric_values(ws) -> dict[str, object]:
    metrics: dict[str, object] = {}
    for row in ws.iter_rows(values_only=True):
        if len(row) >= 3 and row[1] in {'Opening Equity', 'Closing Equity', 'Difference', 'Net Brokerage'}:
            metrics[str(row[1])] = row[2]
    return metrics

def _build_workbook(clip_negative_equity: bool):
    payload = _base_payload(clip_negative_equity)
    workbook_buffer = api.build_brokerage_export_workbook(payload)
    workbook_buffer.seek(0)
    return load_workbook(workbook_buffer, data_only=True)

class BrokerageNegativeEquityTests(unittest.TestCase):

    def test_manual_fixture_matches_expected_delta(self):
        raw_values = [equity for _, _, equity in NEGATIVE_EQUITY_FIXTURE] + [POSITIVE_OPENING_EQ]
        frame = pd.DataFrame({'Equity': raw_values})
        clamped = api._floor_negative_equity_dataframe(frame)
        self.assertAlmostEqual(float(frame['Equity'].sum()), 517.25, places=2)
        self.assertAlmostEqual(float(clamped['Equity'].sum()), 1000.0, places=2)
        self.assertAlmostEqual(float(clamped['Equity'].sum() - frame['Equity'].sum()), 482.75, places=2)
        self.assertTrue((clamped['Equity'] >= 0).all())

    def test_export_workbook_respects_clip_negative_equity_flag(self):
        raw_wb = _build_workbook(False)
        clipped_wb = _build_workbook(True)
        raw_accounts = _load_table_rows(raw_wb['Accounts'])
        clipped_accounts = _load_table_rows(clipped_wb['Accounts'])
        raw_opening = _load_table_rows(raw_wb['Opening Snapshot'])
        clipped_opening = _load_table_rows(clipped_wb['Opening Snapshot'])
        raw_closing = _load_table_rows(raw_wb['Closing Snapshot'])
        clipped_closing = _load_table_rows(clipped_wb['Closing Snapshot'])
        raw_metrics = _load_metric_values(raw_wb['PnL'])
        clipped_metrics = _load_metric_values(clipped_wb['PnL'])
        raw_negative_account = next((row for row in raw_accounts if row['Login'] == NEGATIVE_EQUITY_FIXTURE[0][0]))
        clipped_negative_account = next((row for row in clipped_accounts if row['Login'] == NEGATIVE_EQUITY_FIXTURE[0][0]))
        raw_negative_opening = next((row for row in raw_opening if row['Login'] == NEGATIVE_EQUITY_FIXTURE[0][0]))
        clipped_negative_opening = next((row for row in clipped_opening if row['Login'] == NEGATIVE_EQUITY_FIXTURE[0][0]))
        raw_negative_closing = next((row for row in raw_closing if row['Login'] == NEGATIVE_EQUITY_FIXTURE[0][0]))
        clipped_negative_closing = next((row for row in clipped_closing if row['Login'] == NEGATIVE_EQUITY_FIXTURE[0][0]))
        self.assertEqual(raw_negative_account['Opening Equity'], -210.0)
        self.assertEqual(raw_negative_account['Closing Equity'], -210.0)
        self.assertEqual(raw_negative_opening['Equity'], -210.0)
        self.assertEqual(raw_negative_closing['Equity'], -210.0)
        self.assertEqual(clipped_negative_account['Opening Equity'], 0.0)
        self.assertEqual(clipped_negative_account['Closing Equity'], 0.0)
        self.assertEqual(clipped_negative_opening['Equity'], 0.0)
        self.assertEqual(clipped_negative_closing['Equity'], 0.0)
        self.assertEqual(raw_metrics['Opening Equity'], '$517.25')
        self.assertEqual(raw_metrics['Closing Equity'], '$617.25')
        self.assertEqual(raw_metrics['Difference'], '$100.00')
        self.assertEqual(raw_metrics['Net Brokerage'], '$100.00')
        self.assertEqual(clipped_metrics['Opening Equity'], '$1,000.00')
        self.assertEqual(clipped_metrics['Closing Equity'], '$1,100.00')
        self.assertEqual(clipped_metrics['Difference'], '$100.00')
        self.assertEqual(clipped_metrics['Net Brokerage'], '$100.00')
if __name__ == '__main__':
    unittest.main()
