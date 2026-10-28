"""Persistent UTC history. No reporting function calls MT5.

Money is stored as decimal text and aggregated with Decimal, never SQLite REAL.
Account currencies are never combined. Source rows remain auditable JSON.
"""
from __future__ import annotations

from contextlib import contextmanager
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
import csv
import hashlib
import io
import json
import sqlite3
import threading
import time
import uuid


def utc(ts=None):
    return datetime.fromtimestamp(time.time() if ts is None else ts, timezone.utc).isoformat()


def number(value):
    if value is None:
        return None
    try:
        result = Decimal(str(value))
        return result if result.is_finite() else None
    except (InvalidOperation, ValueError):
        return None


def text(value):
    value = number(value)
    return str(value) if value is not None else None


class DecimalSum:
    def __init__(self):
        self.total, self.count = Decimal(0), 0

    def step(self, value):
        value = number(value)
        if value is not None:
            self.total += value
            self.count += 1

    def finalize(self):
        return str(self.total) if self.count else None


class DecimalMin:
    def __init__(self): self.value = None
    def step(self, value):
        value = number(value)
        if value is not None: self.value = value if self.value is None else min(self.value,value)
    def finalize(self): return text(self.value)


class DecimalMax(DecimalMin):
    def step(self, value):
        value = number(value)
        if value is not None: self.value = value if self.value is None else max(self.value,value)


class DecimalAverage(DecimalSum):
    def finalize(self): return str(self.total/self.count) if self.count else None


SCHEMA = """
CREATE TABLE IF NOT EXISTS history_migrations(version INTEGER PRIMARY KEY, applied_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS history_source(id INTEGER PRIMARY KEY CHECK(id=1), fingerprint TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS history_accounts(login INTEGER PRIMARY KEY,account_group TEXT NOT NULL,currency TEXT NOT NULL,
 name TEXT NOT NULL,updated_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS history_risk_settings(id INTEGER PRIMARY KEY CHECK(id=1),payload TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS history_settings(id INTEGER PRIMARY KEY CHECK(id=1), interval_minutes INTEGER NOT NULL DEFAULT 60,
 snapshot_seconds INTEGER NOT NULL DEFAULT 300, reconciliation_hour INTEGER NOT NULL DEFAULT 2,
 backfill_from TEXT NOT NULL, overlap_days INTEGER NOT NULL DEFAULT 2, reconcile_days INTEGER NOT NULL DEFAULT 30);
CREATE TABLE IF NOT EXISTS history_jobs(id TEXT PRIMARY KEY, kind TEXT NOT NULL, status TEXT NOT NULL,
 requested_at REAL NOT NULL, started_at REAL, completed_at REAL, owner TEXT, lease_until REAL,
 fetched INTEGER NOT NULL DEFAULT 0, inserted INTEGER NOT NULL DEFAULT 0, updated INTEGER NOT NULL DEFAULT 0,
 skipped INTEGER NOT NULL DEFAULT 0, failed_batches INTEGER NOT NULL DEFAULT 0, retries INTEGER NOT NULL DEFAULT 0,
 error TEXT, actor TEXT NOT NULL);
CREATE INDEX IF NOT EXISTS history_jobs_status ON history_jobs(status,requested_at);
CREATE TABLE IF NOT EXISTS history_leases(name TEXT PRIMARY KEY,owner TEXT NOT NULL,expires REAL NOT NULL);
CREATE TABLE IF NOT EXISTS history_checkpoints(account_group TEXT PRIMARY KEY, through_ts INTEGER NOT NULL,
 coverage_from INTEGER NOT NULL, updated_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS history_coverage(account_group TEXT NOT NULL,from_ts INTEGER NOT NULL,to_ts INTEGER NOT NULL,
 PRIMARY KEY(account_group,from_ts));
CREATE TABLE IF NOT EXISTS history_batch_audits(job_id TEXT NOT NULL,account_group TEXT NOT NULL,from_ts INTEGER NOT NULL,to_ts INTEGER NOT NULL,
 source_records INTEGER NOT NULL,stored_records INTEGER NOT NULL,source_net TEXT NOT NULL,stored_net TEXT NOT NULL,status TEXT NOT NULL,checked_at TEXT NOT NULL,
 PRIMARY KEY(job_id,account_group,from_ts,to_ts));
CREATE INDEX IF NOT EXISTS history_batch_audits_checked ON history_batch_audits(checked_at DESC);
CREATE TABLE IF NOT EXISTS history_records(kind TEXT NOT NULL, source_id TEXT NOT NULL, login INTEGER NOT NULL,
 account_group TEXT NOT NULL, currency TEXT NOT NULL, symbol TEXT NOT NULL, time_ms INTEGER NOT NULL,
 action INTEGER, entry INTEGER, profit TEXT, commission TEXT, swap TEXT, fee TEXT, net TEXT, lots TEXT,
 payload TEXT NOT NULL, fingerprint TEXT NOT NULL, updated_at TEXT NOT NULL,
 PRIMARY KEY(kind,source_id));
CREATE INDEX IF NOT EXISTS history_records_scope ON history_records(kind,account_group,time_ms,login,symbol);
CREATE INDEX IF NOT EXISTS history_records_time ON history_records(kind,time_ms);
CREATE TABLE IF NOT EXISTS history_daily(day TEXT NOT NULL,account_group TEXT NOT NULL,login INTEGER NOT NULL,
 currency TEXT NOT NULL,symbol TEXT NOT NULL,profit TEXT,commission TEXT,swap TEXT,fee TEXT,net TEXT,lots TEXT,
 deals INTEGER NOT NULL,buy_deals INTEGER NOT NULL,sell_deals INTEGER NOT NULL,
 PRIMARY KEY(day,account_group,login,currency,symbol));
CREATE INDEX IF NOT EXISTS history_daily_scope ON history_daily(day,account_group,login,symbol,currency);
CREATE TABLE IF NOT EXISTS history_snapshots(source_ts REAL NOT NULL,login INTEGER NOT NULL,account_group TEXT NOT NULL,
 currency TEXT NOT NULL,balance TEXT,equity TEXT,margin TEXT,free_margin TEXT,floating TEXT,positions INTEGER NOT NULL,
 PRIMARY KEY(source_ts,login));
CREATE INDEX IF NOT EXISTS history_snapshots_scope ON history_snapshots(source_ts,account_group,login,currency);
CREATE INDEX IF NOT EXISTS history_snapshots_login ON history_snapshots(login,source_ts DESC);
CREATE TABLE IF NOT EXISTS history_positions(source_ts REAL NOT NULL,ticket INTEGER NOT NULL,login INTEGER NOT NULL,
 account_group TEXT NOT NULL,currency TEXT NOT NULL,symbol TEXT NOT NULL,side TEXT NOT NULL,lots TEXT,
 profit TEXT,contract_size TEXT,price TEXT,rate_profit TEXT,notional TEXT,payload TEXT NOT NULL,
 PRIMARY KEY(source_ts,ticket));
CREATE INDEX IF NOT EXISTS history_positions_scope ON history_positions(source_ts,account_group,login,symbol,currency);
CREATE INDEX IF NOT EXISTS history_positions_login ON history_positions(login,source_ts);
CREATE TABLE IF NOT EXISTS history_symbols(symbol TEXT PRIMARY KEY,source_ts REAL NOT NULL,payload TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS history_discrepancies(id INTEGER PRIMARY KEY,job_id TEXT,kind TEXT,source_id TEXT,
 previous_hash TEXT,new_hash TEXT,detected_at TEXT,detail TEXT);
CREATE TABLE IF NOT EXISTS history_exports(id TEXT PRIMARY KEY,status TEXT NOT NULL,actor TEXT NOT NULL,
 format TEXT NOT NULL,filters TEXT NOT NULL,created_at REAL NOT NULL,completed_at REAL,content BLOB,error TEXT);
"""


class ArchiveSourceMismatch(RuntimeError):
    """A configured Manager must never read or overwrite another source's archive."""


class HistoryStore:
    def __init__(self, path):
        self.path = str(path)
        self._local = threading.local()

    @contextmanager
    def connect(self):
        existing = getattr(self._local,"read_connection",None)
        if existing is not None:
            yield existing
            return
        con = sqlite3.connect(self.path, timeout=15)
        con.row_factory = sqlite3.Row
        con.create_aggregate("decimal_sum", 1, DecimalSum)
        con.create_aggregate("decimal_min", 1, DecimalMin)
        con.create_aggregate("decimal_max", 1, DecimalMax)
        con.create_aggregate("decimal_avg", 1, DecimalAverage)
        con.create_function("decimal_sign",1,lambda v: None if number(v) is None else (number(v)>0)-(number(v)<0))
        con.create_function("decimal_neg",1,lambda v: str(-number(v)) if number(v) is not None else None)
        con.create_collation("DECIMAL",lambda a,b: (Decimal(a)>Decimal(b))-(Decimal(a)<Decimal(b)))
        con.execute("PRAGMA busy_timeout=15000")
        con.execute("PRAGMA synchronous=NORMAL")
        try:
            with con:
                yield con
        finally:
            con.close()

    def migrate(self, backfill_from, interval=60, snapshot_seconds=300, reconciliation_hour=2):
        with self.connect() as con:
            con.executescript(SCHEMA)
            columns = {r[1] for r in con.execute("PRAGMA table_info(history_exports)")}
            for column,kind in (("owner","TEXT"),("lease_until","REAL")):
                if column not in columns:
                    con.execute(f"ALTER TABLE history_exports ADD COLUMN {column} {kind}")
            if not con.execute("SELECT 1 FROM history_migrations WHERE version=2").fetchone():
                con.execute("INSERT OR IGNORE INTO history_coverage SELECT account_group,coverage_from,through_ts FROM history_checkpoints")
                con.execute("INSERT INTO history_migrations VALUES(2,?)",(utc(),))
            con.execute("INSERT OR IGNORE INTO history_migrations VALUES(1,?)", (utc(),))
            con.execute("INSERT OR IGNORE INTO history_settings(id,backfill_from,interval_minutes,snapshot_seconds,reconciliation_hour) VALUES(1,?,?,?,?)",
                        (backfill_from, interval, snapshot_seconds, reconciliation_hour))
            con.execute("INSERT OR IGNORE INTO history_risk_settings VALUES(1,?)",(json.dumps({"enabled":True,"exposure_increase_pct":100,"minimum_lot_increase":1,"floating_change_pct":20,"minimum_floating_change":100,"margin_utilization_pct":80}),))

    def settings(self):
        with self.connect() as con:
            result = dict(con.execute("SELECT * FROM history_settings WHERE id=1").fetchone())
            result["risk_thresholds"] = json.loads(con.execute("SELECT payload FROM history_risk_settings WHERE id=1").fetchone()[0])
            return result

    def bind_source(self, identity):
        fingerprint = hashlib.sha256(identity.encode()).hexdigest()
        with self.connect() as con:
            con.execute("INSERT OR IGNORE INTO history_source VALUES(1,?)",(fingerprint,))
            if con.execute("SELECT fingerprint FROM history_source WHERE id=1").fetchone()[0] != fingerprint:
                raise ArchiveSourceMismatch("Archive belongs to another Manager source; use a separate database")

    def verify_source(self, identity):
        """Read-only guard for reports; an unbound empty archive is still readable."""
        fingerprint = hashlib.sha256(identity.encode()).hexdigest()
        with self.connect() as con:
            row = con.execute("SELECT fingerprint FROM history_source WHERE id=1").fetchone()
            if row and row[0] != fingerprint:
                raise ArchiveSourceMismatch("Archive belongs to another Manager source; use a separate database")

    def configure(self, interval_minutes, snapshot_seconds, reconciliation_hour, backfill_from, risk_thresholds=None):
        with self.connect() as con:
            con.execute("UPDATE history_settings SET interval_minutes=?,snapshot_seconds=?,reconciliation_hour=?,backfill_from=? WHERE id=1",
                        (interval_minutes, snapshot_seconds, reconciliation_hour, backfill_from))
            if risk_thresholds is not None:
                con.execute("UPDATE history_risk_settings SET payload=? WHERE id=1",(json.dumps(risk_thresholds),))

    def lease(self, name, owner, seconds=180):
        now = time.time()
        with self.connect() as con:
            result = con.execute("""INSERT INTO history_leases VALUES(?,?,?) ON CONFLICT(name) DO UPDATE
                SET owner=excluded.owner,expires=excluded.expires WHERE history_leases.expires<? OR history_leases.owner=?""",
                                 (name, owner, now + seconds, now, owner))
            return result.rowcount == 1

    def release(self, name, owner):
        with self.connect() as con:
            con.execute("DELETE FROM history_leases WHERE name=? AND owner=?", (name, owner))

    def enqueue(self, kind="incremental", actor="scheduler"):
        with self.connect() as con:
            con.execute("BEGIN IMMEDIATE")
            existing = con.execute("SELECT id FROM history_jobs WHERE status IN ('queued','running') LIMIT 1").fetchone()
            if existing:
                return existing[0]
            recent = con.execute("SELECT requested_at FROM history_jobs ORDER BY requested_at DESC LIMIT 1").fetchone()
            if recent and time.time() - recent[0] < 30:
                return None
            job_id = uuid.uuid4().hex
            con.execute("INSERT INTO history_jobs(id,kind,status,requested_at,actor) VALUES(?,?,'queued',?,?)",
                        (job_id, kind, time.time(), actor))
            return job_id

    def claim(self, owner):
        with self.connect() as con:
            con.execute("BEGIN IMMEDIATE")
            now = time.time()
            con.execute("UPDATE history_jobs SET status='queued',retries=retries+1,error='Worker lease expired; resuming from checkpoint' WHERE status='running' AND lease_until<?", (now,))
            row = con.execute("SELECT * FROM history_jobs WHERE status='queued' ORDER BY requested_at LIMIT 1").fetchone()
            if not row:
                return None
            con.execute("UPDATE history_jobs SET status='running',owner=?,started_at=?,lease_until=? WHERE id=?",
                        (owner, now, now + 180, row["id"]))
            return dict(row)

    def heartbeat(self, job_id, owner):
        if not self.lease("history-sync", owner):
            raise RuntimeError("Historical worker lease lost")
        with self.connect() as con:
            changed = con.execute("UPDATE history_jobs SET lease_until=? WHERE id=? AND owner=? AND status='running'",
                                  (time.time() + 180, job_id, owner)).rowcount
            if not changed:
                raise RuntimeError("Historical job ownership lost")

    def finish(self, job_id, error=None, owner=None):
        with self.connect() as con:
            con.execute("UPDATE history_jobs SET status=?,completed_at=?,error=? WHERE id=? AND (? IS NULL OR owner=?)",
                        ("failed" if error else "completed", time.time(), error, job_id,owner,owner))

    def checkpoint(self, group):
        with self.connect() as con:
            row = con.execute("SELECT * FROM history_checkpoints WHERE account_group=?", (group,)).fetchone()
            return dict(row) if row else None

    def backfill_start(self, group, initial):
        with self.connect() as con:
            ranges = con.execute("SELECT from_ts,to_ts FROM history_coverage WHERE account_group=? ORDER BY from_ts",(group,)).fetchall()
        cursor = initial
        for start,end in ranges:
            if start > cursor:
                break
            if end >= cursor:
                cursor = end+1
        return cursor

    def unknown_currency_groups(self):
        with self.connect() as con:
            return [r[0] for r in con.execute("SELECT DISTINCT account_group FROM history_records WHERE currency LIKE 'UNKNOWN:%'")]

    def resolve_group_currency(self, group, currency):
        """Re-normalize only unresolved source rows using verified group metadata."""
        with self.connect() as con:
            rows = con.execute("SELECT payload FROM history_records WHERE account_group=?",(group,)).fetchall()
        records=[]
        changed=0
        for row in rows:
            record=json.loads(row[0])
            if record["currency"].startswith("UNKNOWN:"):
                record["currency"]=currency
                changed+=1
            records.append(record)
        if changed:
            start=min(r["time_ms"] for r in records)//1000
            end=max(r["time_ms"] for r in records)//1000
            self.save_batch("metadata-"+uuid.uuid4().hex,group,start,end,records,advance=False)
        return changed

    def _record_coverage(self, con, group, start, end):
        old = con.execute("SELECT * FROM history_checkpoints WHERE account_group=?",(group,)).fetchone()
        ranges = con.execute("SELECT from_ts,to_ts FROM history_coverage WHERE account_group=? AND from_ts<=? AND to_ts>=?",(group,end+1,start-1)).fetchall()
        for left,right in ranges:
            start,end = min(start,left),max(end,right)
            con.execute("DELETE FROM history_coverage WHERE account_group=? AND from_ts=?",(group,left))
        con.execute("INSERT INTO history_coverage VALUES(?,?,?)",(group,start,end))
        # Keep the published checkpoint on a contiguous interval. An earlier,
        # disconnected partial backfill must never advertise coverage over a gap.
        anchor = old["through_ts"] if old else start
        contiguous = con.execute("SELECT from_ts,to_ts FROM history_coverage WHERE account_group=? AND from_ts<=? AND to_ts>=?",(group,anchor,anchor)).fetchone()
        if contiguous:
            con.execute("INSERT OR REPLACE INTO history_checkpoints VALUES(?,?,?,?)",(group,contiguous[1],contiguous[0],utc()))

    def _audit_batch(self, con, job_id, group, start, end, records):
        source = {(r["kind"],r["source_id"]):r for r in records if r["kind"] == "deal" and start*1000 <= r["time_ms"] < (end+1)*1000}
        stored = [dict(r) for r in con.execute("SELECT currency,action,net FROM history_records WHERE kind='deal' AND account_group=? AND time_ms>=? AND time_ms<?",(group,start*1000,(end+1)*1000))]
        def totals(rows):
            result = {}
            for row in rows:
                if row["action"] not in (0,1):
                    continue
                currency = row["currency"]
                if currency not in result: result[currency] = Decimal(0)
                if row.get("net") is None: result[currency] = None
                elif result[currency] is not None: result[currency] += Decimal(row["net"])
            return {key:text(value) for key,value in sorted(result.items())}
        source_net,stored_net = totals(source.values()),totals(stored)
        state = "matched" if len(source) == len(stored) and source_net == stored_net else "investigate"
        con.execute("INSERT OR REPLACE INTO history_batch_audits VALUES(?,?,?,?,?,?,?,?,?,?)",
                    (job_id,group,start,end,len(source),len(stored),json.dumps(source_net),json.dumps(stored_net),state,utc()))

    def save_batch(self, job_id, group, start, end, records, advance=True, owner=None):
        """Atomic records, daily aggregates and checkpoint. Incomplete batches never advance."""
        inserted = updated = skipped = 0
        affected = set()
        with self.connect() as con:
            con.execute("BEGIN IMMEDIATE")
            if owner:
                owned = con.execute("SELECT 1 FROM history_leases WHERE name='history-sync' AND owner=? AND expires>?",(owner,time.time())).fetchone()
                job = con.execute("SELECT 1 FROM history_jobs WHERE id=? AND owner=? AND status='running'",(job_id,owner)).fetchone()
                if not owned or not job:
                    raise RuntimeError("Historical batch lost ownership")
            for row in records:
                payload = json.dumps(row, sort_keys=True, separators=(",", ":"), allow_nan=False)
                fingerprint = hashlib.sha256(payload.encode()).hexdigest()
                old = con.execute("SELECT fingerprint,time_ms,account_group FROM history_records WHERE kind=? AND source_id=?", (row["kind"], row["source_id"])).fetchone()
                if old and old["fingerprint"] == fingerprint:
                    skipped += 1
                    continue
                if old:
                    updated += 1
                    affected.add((datetime.fromtimestamp(old["time_ms"] / 1000, timezone.utc).date().isoformat(),old["account_group"]))
                    con.execute("INSERT INTO history_discrepancies(job_id,kind,source_id,previous_hash,new_hash,detected_at,detail) VALUES(?,?,?,?,?,?,?)",
                                (job_id,row["kind"],row["source_id"],old["fingerprint"],fingerprint,utc(),"Source correction applied"))
                else:
                    inserted += 1
                fields = ["kind","source_id","login","account_group","currency","symbol","time_ms","action","entry","profit","commission","swap","fee","net","lots"]
                values = [row.get(key) for key in fields] + [payload, fingerprint, utc()]
                con.execute(f"INSERT OR REPLACE INTO history_records VALUES({','.join('?' for _ in values)})", values)
                affected.add((datetime.fromtimestamp(row["time_ms"] / 1000, timezone.utc).date().isoformat(),row["account_group"]))
            for day, affected_group in affected:
                con.execute("DELETE FROM history_daily WHERE day=? AND account_group=?", (day,affected_group))
                con.execute("""INSERT INTO history_daily SELECT strftime('%Y-%m-%d',time_ms/1000,'unixepoch'),account_group,login,currency,symbol,
                 CASE WHEN COUNT(profit)=COUNT(*) THEN decimal_sum(profit) END,
                 CASE WHEN COUNT(commission)=COUNT(*) THEN decimal_sum(commission) END,
                 CASE WHEN COUNT(swap)=COUNT(*) THEN decimal_sum(swap) END,
                 CASE WHEN COUNT(fee)=COUNT(*) THEN decimal_sum(fee) END,
                 CASE WHEN COUNT(net)=COUNT(*) THEN decimal_sum(net) END,
                 CASE WHEN COUNT(lots)=COUNT(*) THEN decimal_sum(lots) END,
                 COUNT(*),SUM(action=0),SUM(action=1) FROM history_records WHERE kind='deal' AND account_group=? AND action IN (0,1)
                 AND time_ms>=? AND time_ms<? GROUP BY 1,2,3,4,5""",
                            (affected_group,int(datetime.fromisoformat(day).replace(tzinfo=timezone.utc).timestamp()) * 1000,
                             (int(datetime.fromisoformat(day).replace(tzinfo=timezone.utc).timestamp()) + 86400) * 1000))
            if advance:
                self._record_coverage(con,group,start,end)
            self._audit_batch(con,job_id,group,start,end,records)
            con.execute("UPDATE history_jobs SET fetched=fetched+?,inserted=inserted+?,updated=updated+?,skipped=skipped+? WHERE id=?",
                        (len(records),inserted,updated,skipped,job_id))
        return {"inserted": inserted, "updated": updated, "skipped": skipped}

    def snapshots(self, view, symbols=None):
        source_ts = float(view.get("last_update") or 0)
        if source_ts <= 0 or time.time() - source_ts > 60 or view.get("errors"):
            return False
        users = {int(u["Login"]): u for u in view["users"]}
        positions = {}
        for position in view["positions"]:
            positions.setdefault(int(position["Login"]), []).append(position)
        thresholds = self.settings()["risk_thresholds"]
        with self.connect() as con:
            risk_available = thresholds["enabled"] and con.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='risk_events'").fetchone()
            for login, user in users.items():
                group, currency = user.get("Group", ""), user.get("Currency") or f"UNKNOWN:{login}"
                con.execute("INSERT OR REPLACE INTO history_accounts VALUES(?,?,?,?,?)",(login,group,currency,str(user.get("Name") or ""),utc(source_ts)))
                rows = positions.get(login, [])
                previous = con.execute("SELECT * FROM history_snapshots WHERE login=? AND source_ts<? ORDER BY source_ts DESC LIMIT 1",(login,source_ts)).fetchone() if risk_available else None
                con.execute("INSERT OR IGNORE INTO history_snapshots VALUES(?,?,?,?,?,?,?,?,?,?)",
                            (source_ts,login,group,currency,text(user.get("Balance")),text(user.get("Equity")),text(user.get("Margin")),
                             text(user.get("Free Margin")),text(user.get("Floating PnL")),len(rows)))
                for p in rows:
                    lots, contract, price, rate = [number(p.get(key)) for key in ("Volume","ContractSize","PriceCurrent","RateProfit")]
                    # Price * contract is only valid for linear forex/CFD calculation modes.
                    spec = (symbols or {}).get(p.get("Symbol"), {})
                    mode = spec.get("CalcMode")
                    notional = abs(lots*contract*price*rate) if mode in (0,2,4,5,32,38) and all(v is not None and v > 0 for v in (lots,contract,price,rate)) else None
                    con.execute("INSERT OR IGNORE INTO history_positions VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                                (source_ts,int(p.get("Ticket") or 0),login,group,currency,p.get("Symbol", ""),
                                 "BUY" if p.get("Action") == 0 else "SELL" if p.get("Action") == 1 else "UNKNOWN",
                                 text(lots),text(p.get("Profit")),text(contract),text(price),text(rate),text(notional),json.dumps(p,default=str)))
                if risk_available:
                    self._snapshot_signals(con,source_ts,login,group,currency,user,rows,previous,thresholds)
            for symbol, spec in (symbols or {}).items():
                con.execute("INSERT OR REPLACE INTO history_symbols VALUES(?,?,?)", (symbol,source_ts,json.dumps(spec,default=str)))
        return True

    def _snapshot_signals(self, con, source_ts, login, group, currency, user, positions, previous, thresholds):
        def signal(detector, actual, evidence, symbol="", severity="high"):
            event_id = hashlib.sha256(f"archive|{detector}|{login}|{symbol}|{source_ts}".encode()).hexdigest()
            stamp = utc(source_ts)
            payload = {"evidence":evidence,"observed_at":stamp,"previous_observed_at":utc(previous["source_ts"]) if previous else None,"actual_values":actual}
            con.execute("""INSERT OR IGNORE INTO risk_events(event_id,event_type,detector,login,symbol,account_group,severity,confidence,status,source,
             source_generation,detected_at,first_seen_at,last_seen_at,evidence_json,thresholds_json,actual_values_json,data_quality)
             VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                        (event_id,"historical_snapshot",detector,login,symbol,group,severity,"VERIFIED","NEW","observed_mt5_snapshots",str(source_ts),stamp,stamp,stamp,
                         json.dumps(payload),json.dumps(thresholds),json.dumps(actual),"VERIFIED"))
        margin,equity = number(user.get("Margin")),number(user.get("Equity"))
        if margin is not None and equity is not None and equity > 0:
            utilization = margin/equity*100
            if utilization >= Decimal(str(thresholds["margin_utilization_pct"])):
                signal("observed_margin_utilization",{"margin":text(margin),"equity":text(equity),"utilization_pct":text(utilization),"currency":currency},"Observed account margin / equity exceeds the configured threshold.")
        if previous is None or previous["currency"] != currency or previous["account_group"] != group:
            return
        current_float,old_float,old_balance = number(user.get("Floating PnL")),number(previous["floating"]),number(previous["balance"])
        if not currency.startswith("UNKNOWN") and current_float is not None and old_float is not None and old_balance is not None and abs(old_balance) > 0:
            change = abs(current_float-old_float)
            if change >= Decimal(str(thresholds["minimum_floating_change"])) and change/abs(old_balance)*100 >= Decimal(str(thresholds["floating_change_pct"])):
                signal("observed_floating_pnl_change",{"previous_floating":text(old_float),"current_floating":text(current_float),"absolute_change":text(change),"previous_balance":text(old_balance),"currency":currency},"Change between two observed floating P&L snapshots exceeds the configured share of prior balance. Cash flows and market movement require review.")
        prior = {r[0]:number(r[1]) for r in con.execute("SELECT symbol,decimal_sum(lots) FROM history_positions WHERE login=? AND source_ts=? GROUP BY symbol",(login,previous["source_ts"]))}
        current = {}
        for row in positions:
            lots = number(row.get("Volume"))
            symbol = row.get("Symbol","")
            current[symbol] = None if lots is None or (symbol in current and current[symbol] is None) else current.get(symbol,Decimal(0))+lots
        for symbol,lots in current.items():
            old = prior.get(symbol)
            if old is not None and old > 0 and lots is not None:
                change = lots-old
                if change >= Decimal(str(thresholds["minimum_lot_increase"])) and change/old*100 >= Decimal(str(thresholds["exposure_increase_pct"])):
                    signal("observed_symbol_exposure_increase",{"previous_gross_lots":text(old),"current_gross_lots":text(lots),"increase_pct":text(change/old*100)},"Gross lots in this instrument increased between two actual observations. This is a review signal, not a fraud finding.",symbol)

    def status(self):
        with self.connect() as con:
            jobs = [dict(r) for r in con.execute("SELECT * FROM history_jobs ORDER BY requested_at DESC LIMIT 20")]
            coverage = [dict(r) for r in con.execute("SELECT * FROM history_checkpoints ORDER BY account_group")]
            last = con.execute("SELECT MAX(completed_at) FROM history_jobs WHERE status='completed'").fetchone()[0]
            snapshot = con.execute("SELECT MAX(source_ts) FROM history_snapshots").fetchone()[0]
            discrepancies = con.execute("SELECT COUNT(*) FROM history_discrepancies").fetchone()[0]
            last_reconcile = con.execute("SELECT MAX(completed_at) FROM history_jobs WHERE status='completed' AND kind='reconcile'").fetchone()[0]
            audits = [dict(r) for r in con.execute("SELECT * FROM history_batch_audits ORDER BY checked_at DESC LIMIT 20")]
            mismatches = con.execute("SELECT COUNT(*) FROM history_batch_audits WHERE status='investigate'").fetchone()[0]
        for job in jobs:
            job["duration_seconds"] = round((job["completed_at"] or time.time())-job["started_at"],2) if job["started_at"] else None
        settings = self.settings()
        return {"settings":settings,"jobs":jobs,"coverage":coverage,"last_success":last,"last_reconciliation":last_reconcile,
                "next_sync":(last + settings["interval_minutes"]*60) if last else None,
                "latest_snapshot":snapshot,"snapshot_stale":not snapshot or time.time()-snapshot > settings["snapshot_seconds"]*2,
                "discrepancies":discrepancies,"reconciliation_audits":audits,"unresolved_comparisons":mismatches,"timezone":"UTC"}

    def claim_export(self, owner):
        with self.connect() as con:
            con.execute("BEGIN IMMEDIATE")
            now = time.time()
            if not con.execute("SELECT 1 FROM history_leases WHERE name='history-export' AND owner=? AND expires>?",(owner,now)).fetchone():
                return None
            con.execute("DELETE FROM history_exports WHERE created_at<? AND status IN ('completed','failed')",(now-86400,))
            con.execute("UPDATE history_exports SET status='queued' WHERE status='running' AND (lease_until IS NULL OR lease_until<?)",(now,))
            row = con.execute("SELECT * FROM history_exports WHERE status='queued' ORDER BY created_at LIMIT 1").fetchone()
            if not row: return None
            con.execute("UPDATE history_exports SET status='running',owner=?,lease_until=? WHERE id=?",(owner,now+180,row["id"]))
            return dict(row)

    def renew_export(self, export_id, owner):
        if not self.lease("history-export",owner):
            return False
        with self.connect() as con:
            return con.execute("UPDATE history_exports SET lease_until=? WHERE id=? AND owner=? AND status='running'",(time.time()+180,export_id,owner)).rowcount == 1

    def finish_export(self, export_id, owner, content=None, error=None):
        with self.connect() as con:
            con.execute("BEGIN IMMEDIATE")
            if not con.execute("SELECT 1 FROM history_leases WHERE name='history-export' AND owner=? AND expires>?",(owner,time.time())).fetchone():
                return False
            return con.execute("UPDATE history_exports SET status=?,completed_at=?,content=?,error=? WHERE id=? AND owner=? AND status='running'",
                               ("failed" if error else "completed",time.time(),content,error,export_id,owner)).rowcount == 1

    @staticmethod
    def where(filters, time_field="day", alias=""):
        clauses, params = [], []
        prefix = f"{alias}." if alias else ""
        if time_field == "day":
            clauses += [f"{prefix}day>=?", f"{prefix}day<=?"]
            params += [filters["from_date"], filters["to_date"]]
        else:
            start = datetime.fromisoformat(filters["from_date"]).replace(tzinfo=timezone.utc).timestamp()
            end = datetime.fromisoformat(filters["to_date"]).replace(tzinfo=timezone.utc).timestamp() + 86400
            scale = 1000 if time_field == "time_ms" else 1
            clauses += [f"{prefix}{time_field}>=?",f"{prefix}{time_field}<?"]
            params += [start*scale,end*scale]
        for key, column in (("group","account_group"),("login","login"),("symbol","symbol"),("currency","currency")):
            if filters.get(key):
                clauses.append(f"{prefix}{column}=?")
                params.append(filters[key])
        if filters.get("groups"):
            placeholders = ','.join('?' for _ in filters["groups"])
            clauses.append(f"{prefix}account_group IN ({placeholders})")
            params.extend(filters["groups"])
        return " AND ".join(clauses), params

    def report(self, filters, dimension="day", page=1, page_size=100, descending=False, sort="label"):
        if dimension not in {"day","week","month","login","account_group","symbol"}:
            raise ValueError("Unknown report dimension")
        where, params = self.where(filters)
        if sort not in {"label","net","lots","deals"}:
            raise ValueError("Unknown report sort")
        expr = {"week":"strftime('%Y-%W',day)","month":"substr(day,1,7)"}.get(dimension, dimension)
        sums = ','.join(f'CASE WHEN COUNT({field})=COUNT(*) THEN decimal_sum({field}) END AS {field}' for field in ('profit','commission','swap','fee','net','lots'))
        query = f"""SELECT {expr} AS label,currency,{sums},
         SUM(deals) AS deals,SUM(buy_deals) AS buy_deals,SUM(sell_deals) AS sell_deals FROM history_daily WHERE {where} GROUP BY 1,2"""
        with self.connect() as con:
            if not con.in_transaction:
                con.execute("BEGIN")
            total = con.execute(f"SELECT COUNT(*) FROM ({query})", params).fetchone()[0]
            sort_sql = f'{sort} COLLATE DECIMAL' if sort in {"net","lots"} else sort
            rows = [dict(r) for r in con.execute(query + f" ORDER BY {sort_sql} {'DESC' if descending else 'ASC'},label,currency LIMIT ? OFFSET ?", (*params,page_size,(page-1)*page_size))]
            totals = [dict(r) for r in con.execute(f"""SELECT currency,{sums},SUM(deals) deals,
             COUNT(DISTINCT login) active_accounts FROM history_daily WHERE {where} GROUP BY currency""", params)]
            stats = [dict(r) for r in con.execute(f"""SELECT currency,SUM(decimal_sign(net)>0) winning_accounts,SUM(decimal_sign(net)<0) losing_accounts,
             SUM(net IS NULL) unavailable_accounts,decimal_min(net) minimum_account_net,decimal_max(net) maximum_account_net,decimal_avg(net) average_account_net
             FROM (SELECT currency,login,CASE WHEN COUNT(net)=COUNT(*) THEN decimal_sum(net) END AS net FROM history_daily WHERE {where} GROUP BY currency,login) GROUP BY currency""",params)]
            if dimension == "login":
                labels = [row["label"] for row in rows]
                metadata = {r[0]:r[1] for r in con.execute(f"SELECT login,name FROM history_accounts WHERE login IN ({','.join('?' for _ in labels)})",labels)} if labels else {}
                for row in rows:
                    row["name"] = metadata.get(row["label"])
        return {"rows":rows,"total":total,"totals":totals,"statistics":stats,"page":page,"page_size":page_size,"timezone":"UTC","generated_at":utc(),
                "definition":"Client net trading P&L = deal Profit + Commission + Storage + Fee on BUY/SELL deals, including entry costs. Balance operations excluded. Volume counts executions, including both entry and exit.",
                "filters":filters}

    def activity(self, filters, kind="deal", page=1, page_size=100, descending=True):
        where, params = self.where(filters,"time_ms")
        if kind == "closed":
            kind = "deal"
            where += " AND action IN (0,1) AND entry IN (1,2,3)"
        with self.connect() as con:
            total = con.execute(f"SELECT COUNT(*) FROM history_records WHERE kind=? AND {where}", (kind,*params)).fetchone()[0]
            rows = [dict(r) for r in con.execute(f"SELECT * FROM history_records WHERE kind=? AND {where} ORDER BY time_ms {'DESC' if descending else 'ASC'},source_id LIMIT ? OFFSET ?",
                                               (kind,*params,page_size,(page-1)*page_size))]
        for row in rows:
            row["source"] = json.loads(row.pop("payload"))
            row.pop("fingerprint")
        return {"rows":rows,"total":total,"timezone":"UTC","generated_at":utc()}

    def snapshot_report(self, filters, exposure=False, page=1, page_size=100, dimension="symbol", latest_only=False):
        scoped = dict(filters)
        # Symbol-filtered account equity cannot be allocated to an instrument.
        if not exposure:
            scoped.pop("symbol", None)
        where, params = self.where(scoped,"source_ts")
        table = "history_positions" if exposure else "history_snapshots"
        if exposure:
            if dimension not in {"symbol","login","account_group"}:
                raise ValueError("Unknown exposure dimension")
            select = f"""source_ts,{dimension} label,symbol,currency,decimal_sum(lots) gross_lots,
             decimal_sum(CASE WHEN side='BUY' THEN lots END) buy_lots,decimal_sum(CASE WHEN side='SELL' THEN lots END) sell_lots,
             decimal_sum(notional) gross_notional,decimal_sum(CASE WHEN side='SELL' THEN decimal_neg(notional) ELSE notional END) net_notional,
             SUM(notional IS NULL) unvalued_positions,COUNT(*) positions"""
            grouping = f"source_ts,{dimension},symbol,currency" if dimension != "symbol" else "source_ts,symbol,currency"
        else:
            if dimension not in {"symbol","login","account_group"}:
                raise ValueError("Unknown account snapshot dimension")
            sums = ','.join(f'CASE WHEN COUNT({field})=COUNT(*) THEN decimal_sum({field}) END AS {field}' for field in ('balance','equity','margin','free_margin','floating'))
            extra = f",{dimension}" if dimension in {"login","account_group"} else ""
            select = f"source_ts,currency{extra},{sums},SUM(positions) positions,COUNT(*) accounts"
            grouping = f"source_ts,currency{extra}"
        with self.connect() as con:
            if not con.in_transaction: con.execute("BEGIN")
            if latest_only:
                latest = con.execute(f"SELECT MAX(source_ts) FROM {table} WHERE {where}",params).fetchone()[0]
                where += " AND source_ts=?"
                params.append(latest)
            query = f"SELECT {select} FROM {table} WHERE {where} GROUP BY {grouping}"
            total = con.execute(f"SELECT COUNT(*) FROM ({query})",params).fetchone()[0]
            rows = [dict(r) for r in con.execute(query + " ORDER BY source_ts DESC LIMIT ? OFFSET ?",(*params,page_size,(page-1)*page_size))]
            values = {}
            if exposure:
                values = {(r[0],r[1]):number(r[2]) for r in con.execute(f"SELECT source_ts,currency,CASE WHEN COUNT(notional)=COUNT(*) THEN decimal_sum(notional) END FROM history_positions WHERE {where} GROUP BY source_ts,currency",params)}
        for row in rows:
            if exposure:
                row["buy_lots"] = row["buy_lots"] or "0"
                row["sell_lots"] = row["sell_lots"] or "0"
                row["net_lots"] = str(Decimal(row["buy_lots"])-Decimal(row["sell_lots"]))
                if row["unvalued_positions"]:
                    row["gross_notional"] = row["net_notional"] = None
                denominator = values.get((row["source_ts"],row["currency"]))
                numerator = number(row["gross_notional"])
                row["concentration_pct"] = text(numerator/denominator*100) if denominator is not None and denominator > 0 and numerator is not None else None
            else:
                margin,equity = number(row["margin"]),number(row["equity"])
                row["margin_utilization_pct"] = text(margin/equity*100) if margin is not None and equity is not None and equity > 0 else None
                row["margin_level_pct"] = text(equity/margin*100) if margin is not None and margin > 0 and equity is not None else None
        return {"rows":rows,"total":total,"timezone":"UTC","generated_at":utc(),"latest_only":latest_only,
                "definition":"Observed snapshots only; gaps are not interpolated. Gross lots = buy + sell; net lots = buy - sell per symbol. Notional requires supported calculation mode, contract size, price and conversion rate. Equity is account-wide and cannot be filtered by symbol."}

    def risk_report(self, filters, page=1, page_size=100, severity=""):
        where = ["first_seen_at>=?","first_seen_at<?"]
        from datetime import timedelta
        end = datetime.fromisoformat(filters["to_date"])+timedelta(days=1)
        params = [filters["from_date"],end.date().isoformat()]
        for key,column in (("group","account_group"),("login","login"),("symbol","symbol")):
            if filters.get(key):
                where.append(f"{column}=?")
                params.append(filters[key])
        if filters.get("currency"):
            where.append("login IN (SELECT login FROM history_accounts WHERE currency=?)")
            params.append(filters["currency"])
        if severity:
            where.append("severity=?")
            params.append(severity)
        clause = " AND ".join(where)
        with self.connect() as con:
            total = con.execute(f"SELECT COUNT(*) FROM risk_events WHERE {clause}",params).fetchone()[0]
            rows = [dict(r) for r in con.execute(f"SELECT * FROM risk_events WHERE {clause} ORDER BY first_seen_at DESC,event_id LIMIT ? OFFSET ?",(*params,page_size,(page-1)*page_size))]
        for row in rows:
            row["evidence"] = json.loads(row.pop("evidence_json"))
        return {"rows":rows,"total":total,"timezone":"UTC","generated_at":utc(),"definition":"Persisted detection signals and evidence; not proof of fraud. Thresholds and resolution history use existing Operations controls. Currency filters use archived account metadata."}

    def export(self, filters, fmt):
        # Every page and its totals read the same SQLite snapshot even while a
        # synchronization batch or position snapshot commits in another thread.
        with self.connect() as con:
            con.execute("BEGIN")
            self._local.read_connection = con
            try:
                return self._export(filters,fmt)
            finally:
                del self._local.read_connection

    def _export(self, filters, fmt):
        # Runs in the leased worker, never in the request or browser lifecycle.
        requested = filters.get("report","accounts")
        def fetch_page(page):
            if requested == "trading-activity":
                return self.activity(filters,filters.get("kind","deal"),page,1000)
            if requested == "exposure":
                return self.snapshot_report(filters,True,page,1000)
            if requested == "risk-events":
                return self.risk_report(filters,page,1000)
            return self.report(filters,filters.get("dimension","login") if requested == "pnl" else "login",page,1000)
        report = fetch_page(1)
        def records():
            result, page = report, 1
            while True:
                yield from result["rows"]
                if page*1000 >= result["total"]:
                    return
                page += 1
                result = fetch_page(page)
        rows = records()
        columns = ["label","currency","profit","commission","swap","fee","net","lots","deals","buy_deals","sell_deals"]
        if requested == "trading-activity":
            columns = ["time_ms","source_id","login","account_group","currency","symbol","action","entry","lots","profit","commission","swap","fee","net"]
        elif requested == "exposure":
            columns = ["source_ts","symbol","currency","buy_lots","sell_lots","gross_lots","net_lots","gross_notional","net_notional","unvalued_positions"]
        elif requested == "risk-events":
            columns = ["first_seen_at","login","account_group","symbol","detector","severity","status","evidence"]
        # Spreadsheet cells must never evaluate group/login/symbol source text.
        def safe(value):
            if isinstance(value,(dict,list)):
                value = json.dumps(value,sort_keys=True)
            dangerous = isinstance(value,str) and value.lstrip().startswith(("=","+","-","@")) and number(value) is None
            return "'"+value if dangerous else value
        if fmt == "csv":
            output = io.StringIO(newline="")
            writer = csv.writer(output)
            writer.writerow(["Generated UTC",utc(),"From",filters["from_date"],"To",filters["to_date"]])
            writer.writerow(["Filters",json.dumps(filters,sort_keys=True)])
            writer.writerow(["Client net = Profit + Commission + Storage + Fee; currencies are separate"])
            writer.writerow(columns)
            writer.writerows([[safe(row.get(c)) for c in columns] for row in rows])
            writer.writerow([])
            writer.writerow(["Totals by currency"])
            for total in report.get("totals",[]):
                writer.writerow(["TOTAL",total["currency"],*[total.get(c) for c in columns[2:]]])
            return output.getvalue().encode("utf-8-sig")
        from openpyxl import Workbook
        output = io.BytesIO()
        book = Workbook(write_only=True)
        sheet = book.create_sheet("Risk Analytics")
        sheet.append(["Generated UTC",utc(),"From",filters["from_date"],"To",filters["to_date"]])
        sheet.append(["Filters",json.dumps(filters,sort_keys=True)])
        sheet.append(["Client net = Profit + Commission + Storage + Fee; currencies are separate"])
        sheet.append(columns)
        for row in rows:
            sheet.append([safe(row.get(c)) for c in columns])
        sheet.append(["Totals by currency"])
        for total in report.get("totals",[]):
            sheet.append(["TOTAL",total["currency"],*[total.get(c) for c in columns[2:]]])
        book.save(output)
        return output.getvalue()
